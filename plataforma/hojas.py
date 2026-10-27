#!/usr/bin/env python3
"""Genera el .docx y el .pdf de una sesión a partir de su Markdown de la web.

    uv run -p 3.12 --with python-docx --with pypdf --with reportlab \
        python3 plataforma/hojas.py docs/sesiones/2026-11-03-PPS.md PPS/practicas/PPS-pytest-03-11-2026

Crea <destino>.docx y <destino>.pdf con aspecto de documento de Word normal
(Calibri 11, comandos en Courier sangrados, páginas numeradas). El PDF se imprime
con Chrome en modo headless.

Markdown admitido: # título, ## y ### apartados, párrafos, listas (- y 1.),
tablas, bloques ``` de código, > notas, `código`, **negrita** y [enlaces](url).
Marcas para el papel (en la web no se ven):
    <!-- lineas: 4 -->       líneas para responder a mano
    <!-- pdf: texto -->      párrafo que solo sale en el papel
    <!-- solo-web --> ... <!-- /solo-web -->   bloque que no sale en el papel
"""
from __future__ import annotations

import html
import io
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CAB = "CIPFP Cheste · 1º CE Ciberseguridad en Entornos de las TI · Curso 2026-27"
MODULOS = {
    "PPS": "Puesta en Producción Segura (5023) · Prof. José Picón",
    "IC": "Incidentes de Ciberseguridad (5021) · Prof. José Picón",
    "HE": "Hacking Ético (5025) · Prof. José Picón",
}
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# ---------------------------------------------------------------- Markdown → items
ENLACE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def limpiar_enlaces(texto: str) -> str:
    """En papel, un enlace relativo es solo su texto; uno absoluto lleva la URL."""
    def sub(m: re.Match) -> str:
        txt, url = m.group(1), m.group(2)
        return f"{txt} ({url})" if re.match(r"^https?://", url) and url not in txt else txt
    return ENLACE.sub(sub, texto)


def parse(md: str) -> list[tuple[str, object]]:
    md = re.sub(r"<!--\s*solo-web\s*-->.*?<!--\s*/solo-web\s*-->", "", md, flags=re.S)
    items: list[tuple[str, object]] = []
    lineas = md.splitlines()
    i = 0
    parrafo: list[str] = []

    def cerrar():
        if parrafo:
            items.append(("p", limpiar_enlaces(" ".join(s.strip() for s in parrafo))))
            parrafo.clear()

    while i < len(lineas):
        ln = lineas[i]
        s = ln.strip()
        if s.startswith("```"):
            cerrar()
            bloque = []
            i += 1
            inicio = i
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            if i == len(lineas):
                raise SystemExit(f"Bloque de código sin cerrar (abre en la línea {inicio})")
            items.append(("code", "\n".join(bloque)))
        elif m := re.match(r"<!-- lineas: (\d+) -->", s):
            cerrar()
            items.append(("lines", int(m.group(1))))
        elif m := re.match(r"<!-- pdf: (.*) -->", s):
            cerrar()
            items.append(("p", m.group(1)))
        elif s.startswith("<"):
            cerrar()  # HTML de la web (divs, comentarios): no va al papel
        elif s.startswith("# "):
            cerrar()
            items.append(("title", s[2:]))
        elif s.startswith("## "):
            cerrar()
            items.append(("h", s[3:]))
        elif s.startswith("### "):
            cerrar()
            items.append(("h3", s[4:]))
        elif s.startswith("|"):
            cerrar()
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [limpiar_enlaces(c.strip()) for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            if filas:
                items.append(("table", filas))
            continue
        elif re.match(r"^(\d+)\. ", s) or s.startswith("- "):
            cerrar()
            tipo = "ol" if s[0].isdigit() else "ul"
            lista: list[str] = []
            while i < len(lineas):
                t = lineas[i].strip()
                if tipo == "ol" and re.match(r"^\d+\. ", t):
                    lista.append(re.sub(r"^\d+\. ", "", t))
                elif tipo == "ul" and t.startswith("- "):
                    lista.append(t[2:])
                elif t and lineas[i].startswith("   ") and lista:
                    lista[-1] += " " + t  # continuación de la línea anterior
                else:
                    break
                i += 1
            items.append((tipo, [limpiar_enlaces(x) for x in lista]))
            continue
        elif s.startswith("> "):
            cerrar()
            nota = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                nota.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            items.append(("p", limpiar_enlaces(" ".join(nota))))
            continue
        elif not s:
            cerrar()
        else:
            parrafo.append(ln)
        i += 1
    cerrar()
    return items


# ---------------------------------------------------------------- en línea
TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def runs(text: str):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("`"):
            yield part[1:-1], "code"
        elif part.startswith("**"):
            yield part[2:-2], "bold"
        else:
            yield part, None


def fuente_courier(r) -> None:
    r.font.name = "Courier New"
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Courier New")
    r.font.size = Pt(10)


def add_runs(par, text: str) -> None:
    for t, kind in runs(text):
        r = par.add_run(t)
        if kind == "code":
            fuente_courier(r)
        elif kind == "bold":
            r.bold = True


# ---------------------------------------------------------------- DOCX
def build_docx(items, path: Path, mod: str) -> None:
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    hp = sec.header.paragraphs[0]
    hp.text = f"{CAB}\n{mod}"
    hp.runs[0].font.size = Pt(8.5)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("José Picón · página ")
    r.font.size = Pt(8.5)
    for fld in ("begin", None, "end"):  # campo PAGE de Word
        r = fp.add_run()
        r.font.size = Pt(8.5)
        if fld:
            el = r._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): fld})
        else:
            el = r._element.makeelement(qn("w:instrText"), {})
            el.text = "PAGE"
        r._element.append(el)
    for kind, val in items:
        if kind == "title":
            p = d.add_paragraph()
            r = p.add_run(val)
            r.bold = True
            r.font.size = Pt(15)
            p.paragraph_format.space_after = Pt(10)
        elif kind in ("h", "h3"):
            p = d.add_paragraph()
            r = p.add_run(val)
            r.bold = True
            r.font.size = Pt(12 if kind == "h" else 11)
            p.paragraph_format.space_before = Pt(10 if kind == "h" else 6)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
        elif kind == "p":
            add_runs(d.add_paragraph(), val)
        elif kind == "code":
            filas = val.split("\n")
            for n, line in enumerate(filas):
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_after = Pt(0) if n < len(filas) - 1 else Pt(6)
                p.paragraph_format.keep_with_next = n < len(filas) - 1
                fuente_courier(p.add_run(line))
        elif kind in ("ol", "ul"):
            for li in val:
                add_runs(d.add_paragraph(style="List Number" if kind == "ol" else "List Bullet"), li)
        elif kind == "lines":
            for _ in range(val):
                p = d.add_paragraph("_" * 86)
                p.paragraph_format.space_after = Pt(4)
        elif kind == "table":
            t = d.add_table(rows=0, cols=len(val[0]))
            t.style = "Table Grid"
            for n, row in enumerate(val):
                cells = t.add_row().cells
                for c, txt in zip(cells, row):
                    c.text = ""
                    for tt, k in runs(txt):
                        r = c.paragraphs[0].add_run(tt)
                        r.font.size = Pt(9.5)
                        r.bold = n == 0 or k == "bold"
                        if k == "code":
                            r.font.name = "Courier New"
            d.add_paragraph()
    d.save(path)


# ---------------------------------------------------------------- HTML → PDF
CSS = """
@page { size: A4; margin: 16mm 20mm 18mm; }
body { font-family: Calibri, Carlito, "Helvetica Neue", Arial, sans-serif; font-size: 11pt;
  line-height: 1.3; color: #000; margin: 0; }
.cab { font-size: 8.5pt; color: #444; border-bottom: 0.5pt solid #888; padding-bottom: 3px; margin-bottom: 14px; }
.title { font-size: 15pt; font-weight: bold; margin: 0 0 10pt; }
h2 { font-size: 12pt; margin: 12pt 0 4pt; break-after: avoid; }
h3 { font-size: 11pt; margin: 8pt 0 3pt; break-after: avoid; }
p { margin: 0 0 6pt; }
code { font-family: "Courier New", Courier, monospace; font-size: 10pt; }
pre { font-family: "Courier New", Courier, monospace; font-size: 9.5pt; margin: 0 0 7pt 1cm;
  white-space: pre-wrap; overflow-wrap: anywhere; break-inside: avoid; }
ol, ul { margin: 0 0 6pt; padding-left: 0.9cm; } li { margin-bottom: 3pt; }
table { border-collapse: collapse; width: 100%; margin: 2pt 0 10pt; font-size: 9.5pt; break-inside: avoid; }
td, th { border: 0.5pt solid #000; padding: 2pt 4pt; vertical-align: top; text-align: left; overflow-wrap: anywhere; }
th { font-weight: bold; }
.line { border-bottom: 0.5pt solid #000; height: 18pt; }
"""


def inline_html(text: str) -> str:
    out = []
    for t, kind in runs(text):
        t = html.escape(t)
        out.append(f"<code>{t}</code>" if kind == "code" else f"<b>{t}</b>" if kind == "bold" else t)
    return "".join(out)


def build_html(items, mod: str) -> str:
    titulo = next((v for k, v in items if k == "title"), "")
    b = [f"<!doctype html><html lang='es'><head><meta charset='utf-8'><title>{html.escape(titulo)}</title>"
         f"<style>{CSS}</style></head><body><div class='cab'>{html.escape(CAB)}<br>{html.escape(mod)}</div>"]
    for kind, val in items:
        if kind == "title":
            b.append(f"<div class='title'>{html.escape(val)}</div>")
        elif kind == "h":
            b.append(f"<h2>{html.escape(val)}</h2>")
        elif kind == "h3":
            b.append(f"<h3>{html.escape(val)}</h3>")
        elif kind == "p":
            b.append(f"<p>{inline_html(val)}</p>")
        elif kind == "code":
            b.append(f"<pre>{html.escape(val)}</pre>")
        elif kind in ("ol", "ul"):
            b.append(f"<{kind}>" + "".join(f"<li>{inline_html(li)}</li>" for li in val) + f"</{kind}>")
        elif kind == "lines":
            b.append("<div class='line'></div>" * val)
        elif kind == "table":
            rows = [f"<tr>{''.join(f'<th>{inline_html(c)}</th>' for c in val[0])}</tr>"]
            rows += [f"<tr>{''.join(f'<td>{inline_html(c)}</td>' for c in r)}</tr>" for r in val[1:]]
            b.append(f"<table>{''.join(rows)}</table>")
    b.append("</body></html>")
    return "".join(b)


def numerar(pdf: Path) -> int:
    """Estampa «José Picón · página N de M» al pie de cada página."""
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas

    lector = PdfReader(str(pdf))
    total = len(lector.pages)
    escritor = PdfWriter()
    for n, pagina in enumerate(lector.pages, 1):
        ancho, alto = float(pagina.mediabox.width), float(pagina.mediabox.height)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(ancho, alto))
        c.setFont("Helvetica", 8)
        c.setFillGray(0.27)
        c.drawRightString(ancho - 57, 28, f"José Picón · página {n} de {total}")
        c.save()
        buf.seek(0)
        pagina.merge_page(PdfReader(buf).pages[0])
        escritor.add_page(pagina)
    with pdf.open("wb") as f:
        escritor.write(f)
    return total


def build_pdf(items, path: Path, mod: str) -> int:
    with tempfile.TemporaryDirectory() as tmp:
        h = Path(tmp) / "hoja.html"
        h.write_text(build_html(items, mod), encoding="utf-8")
        subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={path}", h.as_uri()],
                       check=True, capture_output=True)
    return numerar(path)


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    md, destino = Path(sys.argv[1]), Path(sys.argv[2])
    clave = next((k for k in MODULOS if f"-{k}" in md.stem or md.stem.startswith(k)), None)
    if not clave:
        print(f"No sé de qué módulo es {md.name}")
        return 2
    items = parse(md.read_text(encoding="utf-8"))
    destino.parent.mkdir(parents=True, exist_ok=True)
    build_docx(items, destino.with_suffix(".docx"), MODULOS[clave])
    paginas = build_pdf(items, destino.with_suffix(".pdf"), MODULOS[clave])
    print(f"ok {destino}.docx  {destino}.pdf ({paginas} páginas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
